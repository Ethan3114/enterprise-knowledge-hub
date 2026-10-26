from pathlib import Path

import fitz
import pytest
from docx import Document

from app.chunking import chunk_blocks
from app.limits import LimitError
from app.parsing import ParseError, parse_file


def _settings(**overrides):
    values = {
        "ocr_enabled": True,
        "table_recognition": True,
        "image_text_recognition": True,
        "recognize": None,
    }
    values.update(overrides)
    return values


def test_markdown_keeps_heading_location(tmp_path: Path):
    path = tmp_path / "handbook.md"
    path.write_text("# 年假\n\n员工享有十个工作日年假。\n", encoding="utf-8")
    blocks = parse_file(path, **_settings())
    assert blocks[0].location == "年假"
    assert "十个工作日" in blocks[0].text


def test_csv_uses_header_and_row_limit(tmp_path: Path):
    path = tmp_path / "people.csv"
    path.write_text("姓名,天数\n张三,10\n李四,5\n", encoding="utf-8")
    blocks = parse_file(path, **_settings(), max_table_rows=2)
    assert blocks[0].location == "第2行"
    assert "姓名：张三" in blocks[0].text
    with pytest.raises(LimitError, match="超过 1 行"):
        parse_file(path, **_settings(), max_table_rows=1)


def test_docx_keeps_paragraph_and_table(tmp_path: Path):
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_paragraph("差旅报销需在三十日内提交")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "标准"
    table.rows[1].cells[0].text = "住宿"
    table.rows[1].cells[1].text = "500"
    document.save(path)
    blocks = parse_file(path, **_settings())
    assert any("三十日" in block.text for block in blocks)
    assert any("项目：住宿" in block.text or "住宿" in block.text for block in blocks)


def test_pdf_extracts_text_without_vision(tmp_path: Path):
    path = tmp_path / "leave.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "员工年假为十个工作日，需提前三个工作日申请。", fontname="china-s", fontsize=14)
    pdf.save(path)
    pdf.close()
    blocks = parse_file(path, **_settings(recognize=None))
    assert blocks[0].location == "第1页"
    assert "十个工作日" in blocks[0].text


def test_image_uses_vision_model(tmp_path: Path):
    path = tmp_path / "scan.png"
    pdf = fitz.open()
    page = pdf.new_page(width=240, height=80)
    page.insert_text((20, 40), "年假")
    page.get_pixmap().save(path)
    pdf.close()
    blocks = parse_file(path, **_settings(recognize=lambda _image, _mime: "员工年假十天"))
    assert blocks[0].text == "员工年假十天"
    assert blocks[0].location == "图片"


def test_image_without_recognition_is_rejected(tmp_path: Path):
    path = tmp_path / "scan.png"
    path.write_bytes(b"not-an-image")
    with pytest.raises(ParseError, match="关闭图片文字识别"):
        parse_file(path, **_settings(image_text_recognition=False))


def test_chunk_overlap_is_rejected():
    from app.parsing import TextBlock

    with pytest.raises(ValueError):
        chunk_blocks([TextBlock("年假制度", "全文")], 100, 100)
