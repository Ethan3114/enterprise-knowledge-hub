"""把上传文件变成带位置的文本块。

文字型 PDF、Word、Markdown 和表格在本机提取。
图片和几乎没有文字的扫描页交给已经配置的视觉模型。
"""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

from .limits import MAX_PDF_PAGES, MAX_TABLE_ROWS, LimitError, extension_of

VisionRecognizer = Callable[[bytes, str], str]


@dataclass(frozen=True)
class TextBlock:
    text: str
    location: str


class ParseError(Exception):
    pass


def parse_file(
    path: Path,
    *,
    ocr_enabled: bool,
    table_recognition: bool,
    image_text_recognition: bool,
    recognize: VisionRecognizer | None,
    max_pdf_pages: int = MAX_PDF_PAGES,
    max_table_rows: int = MAX_TABLE_ROWS,
) -> list[TextBlock]:
    extension = extension_of(path.name)
    if extension == "md":
        return _markdown(path)
    if extension == "csv":
        return _csv(path, max_table_rows)
    if extension == "xlsx":
        return _xlsx(path, max_table_rows)
    if extension == "docx":
        return _docx(path, max_table_rows)
    if extension == "pdf":
        return _pdf(
            path,
            ocr_enabled=ocr_enabled,
            table_recognition=table_recognition,
            recognize=recognize,
            max_pdf_pages=max_pdf_pages,
        )
    if extension in {"png", "jpg", "jpeg", "webp", "tif", "tiff"}:
        return _image(path, image_text_recognition, recognize)
    raise ParseError(f"不支持 .{extension or '未知'} 文件")


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("文件编码无法识别，请另存为 UTF-8")


def _markdown(path: Path) -> list[TextBlock]:
    text = _read_text(path).strip()
    if not text:
        return []
    matches = list(re.finditer(r"(?m)^#{1,6} .+$", text))
    if not matches:
        return [TextBlock(text, "全文")]
    blocks: list[TextBlock] = []
    preamble = text[: matches[0].start()].strip()
    if preamble:
        blocks.append(TextBlock(preamble, "全文"))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[match.start() : end].strip()
        title = match.group().lstrip("#").strip() or "全文"
        if body:
            blocks.append(TextBlock(body, title))
    return blocks


def _rows_to_blocks(rows: list[list[str]], location_prefix: str, max_rows: int) -> list[TextBlock]:
    usable = [row for row in rows if any(cell.strip() for cell in row)]
    if not usable:
        return []
    header = usable[0]
    has_header = any(header)
    data_rows = usable[1:] if has_header else usable
    if len(data_rows) > max_rows:
        raise LimitError(f"表格超过 {max_rows} 行")
    start = 2 if has_header else 1
    blocks: list[TextBlock] = []
    for offset, row in enumerate(data_rows):
        pairs: list[str] = []
        for index, cell in enumerate(row):
            value = cell.strip()
            if not value:
                continue
            key = header[index].strip() if has_header and index < len(header) and header[index].strip() else f"列{index + 1}"
            pairs.append(f"{key}：{value}")
        if pairs:
            blocks.append(TextBlock("；".join(pairs), f"{location_prefix}第{start + offset}行"))
    return blocks


def _csv(path: Path, max_rows: int) -> list[TextBlock]:
    text = _read_text(path)
    rows = [[str(cell) for cell in row] for row in csv.reader(io.StringIO(text))]
    return _rows_to_blocks(rows, "", max_rows)


def _xlsx(path: Path, max_rows: int) -> list[TextBlock]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[TextBlock] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            prefix = f"{sheet.title} " if len(workbook.worksheets) > 1 else ""
            blocks.extend(_rows_to_blocks(rows, prefix, max_rows))
    finally:
        workbook.close()
    return blocks


def _docx(path: Path, max_rows: int) -> list[TextBlock]:
    document = Document(str(path))
    blocks: list[TextBlock] = []
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    if paragraphs:
        blocks.append(TextBlock("\n".join(paragraphs), "正文"))
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.extend(_rows_to_blocks(rows, f"表格{index} ", max_rows))
    return blocks


def _pdf(
    path: Path,
    *,
    ocr_enabled: bool,
    table_recognition: bool,
    recognize: VisionRecognizer | None,
    max_pdf_pages: int,
) -> list[TextBlock]:
    document = fitz.open(path)
    try:
        if document.needs_pass:
            raise ParseError("PDF 已加密，请先取消密码再上传")
        if document.page_count > max_pdf_pages:
            raise LimitError(f"PDF 超过 {max_pdf_pages} 页")
        blocks: list[TextBlock] = []
        skipped_scan = False
        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if table_recognition:
                text = _append_tables(page, text)
            location = f"第{index}页"
            # 只有几乎抽不出文字的页面才交给视觉模型，避免把短标题页当成扫描件。
            if len(text) >= 8:
                blocks.append(TextBlock(text, location))
                continue
            if not ocr_enabled:
                if text:
                    blocks.append(TextBlock(text, location))
                else:
                    skipped_scan = True
                continue
            if recognize is None:
                raise ParseError("这一页像扫描件。请管理员先选择视觉识别模型")
            image = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False).tobytes("png")
            recognized = recognize(image, "image/png").strip()
            if recognized:
                blocks.append(TextBlock(recognized, location))
        if not blocks and skipped_scan:
            raise ParseError("没有提取到文字。扫描页需要打开“扫描页识别”并选择视觉模型")
        return blocks
    finally:
        document.close()


def _append_tables(page: fitz.Page, text: str) -> str:
    try:
        finder = page.find_tables()
    except Exception:
        return text
    tables = []
    for table in finder.tables:
        rendered = (table.to_markdown() or "").strip()
        if rendered:
            tables.append(rendered)
    if not tables:
        return text
    return "\n\n".join(part for part in [text, *tables] if part).strip()


def _image(path: Path, enabled: bool, recognize: VisionRecognizer | None) -> list[TextBlock]:
    if not enabled:
        raise ParseError("已关闭图片文字识别，这张图片不会进入索引")
    if recognize is None:
        raise ParseError("请管理员先选择视觉识别模型")
    png = _to_png(path)
    recognized = recognize(png, "image/png").strip()
    if not recognized:
        raise ParseError("视觉模型没有从图片里识别出文字")
    return [TextBlock(recognized, "图片")]


def _to_png(path: Path) -> bytes:
    document = fitz.open(path)
    try:
        if document.page_count < 1:
            raise ParseError("图片无法读取")
        return document[0].get_pixmap(alpha=False).tobytes("png")
    finally:
        document.close()
